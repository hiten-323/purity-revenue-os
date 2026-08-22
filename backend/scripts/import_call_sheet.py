"""
Read a filled call sheet back into Revenue OS.

This is the half that stops the workbook becoming a second CRM. Without it the
spreadsheet accumulates commercial truth the system never sees, and within a
month the sheet outranks the database while nothing enforces either.

    export  ->  call  ->  THIS  ->  Revenue OS  ->  evaluate_next_action()

Rules it will not break
-----------------------
* Trust rises through trust_promoter.on_founder_call(), never by writing
  email_trust directly. That way the shared-inbox cap, the confidence
  recompute and the audit trail all still apply. A sheet must not be able to
  manufacture standing to contact someone.
* Blank never overwrites. An empty cell means "did not ask", not "no".
* A value that will not coerce is reported and skipped, never guessed. If
  someone types "8-10" in Kg/month, that is not the number 9.
* Dry run by default. --commit is required to write anything.

Usage
-----
    python scripts/import_call_sheet.py <file.xlsx>
    python scripts/import_call_sheet.py <file.xlsx> --commit
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from datetime import datetime

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from openpyxl import load_workbook

from call_sheet_schema import COLUMNS, GIVEN

SHEET = "CALL QUEUE"
HEADER_ROW = 3

# The founder heard a yes on a live call. That is first-party and explicit;
# it is not the IMPLIED_B2B that a scraped address gets by default.
CONSENT_YES = "EXPLICIT"
CONSENT_NO = "REFUSED"


def _digits(p) -> str:
    d = re.sub(r"\D", "", str(p or ""))
    return d[-10:] if len(d) >= 10 else d


def _txt(v) -> str:
    s = str(v).strip() if v is not None else ""
    return "" if s.lower() in ("none", "null", "n/a", "nan") else s


def _num(v, cast, field, problems):
    """Coerce or refuse. Never approximate — '8-10 kg' is not 9."""
    s = _txt(v)
    if not s:
        return None
    m = re.fullmatch(r"[^\d.\-]*(\d+(?:\.\d+)?)[^\d]*", s.replace(",", ""))
    if not m:
        problems.append(f"{field}: cannot read a number from {s!r} — skipped")
        return None
    try:
        return cast(m.group(1))
    except ValueError:
        problems.append(f"{field}: {s!r} is not a number — skipped")
        return None


def _date(v, field, problems):
    s = _txt(v)
    if not s:
        return None
    if isinstance(v, datetime):
        return v
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s[:11].strip(), fmt)
        except ValueError:
            continue
    problems.append(f"{field}: cannot read a date from {s!r} — skipped")
    return None


def read_rows(path):
    wb = load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"No '{SHEET}' sheet in {os.path.basename(path)}. "
                         f"Found: {', '.join(wb.sheetnames)}")
    ws = wb[SHEET]

    header = {}
    for i in range(1, ws.max_column + 1):
        t = _txt(ws.cell(HEADER_ROW, i).value)
        if t:
            header[t] = i

    # Match on title, then map back to key, so a reordered sheet still imports
    # and a renamed column fails loudly instead of writing to the wrong field.
    col = {}
    missing = []
    for key, title, _w, group, _dv, target in COLUMNS:
        if title in header:
            col[key] = header[title]
        elif group != GIVEN and target:
            missing.append(title)
    if missing:
        print(f"WARNING: columns not found, will not import: {missing}")

    out = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        row = {k: ws.cell(r, i).value for k, i in col.items()}
        if any(_txt(v) for k, v in row.items()
               if k not in ("row", "company", "segment", "address", "city",
                            "phone", "rating", "reviews", "email_have",
                            "website", "territory", "note")):
            row["_row"] = r
            out.append(row)
    return out


def plan(row, lead, problems):
    """What this row would change. Nothing is written here."""
    ch = {}

    def setf(field, value):
        if value is None or value == "":
            return                      # blank never overwrites
        if getattr(lead, field, None) != value:
            ch[field] = value

    setf("contact_name", _txt(row.get("spoke_with")))
    setf("contact_title", _txt(row.get("their_role")))
    setf("lead_owner", _txt(row.get("owner")))
    setf("call_outcome_last", _txt(row.get("outcome")))
    setf("lead_temperature_tier", _txt(row.get("quality")).upper() or None)
    setf("current_brand", _txt(row.get("brand_now")))
    setf("current_supplier", _txt(row.get("supplier")))
    setf("monthly_consumption", _txt(row.get("kg_month")))
    setf("company_size", _txt(row.get("seats")))
    setf("reorder_sku", _txt(row.get("pack")))
    setf("objection_reason", _txt(row.get("objection")))

    n = _num(row.get("attempt"), int, "Attempt #", problems)
    if n is not None:
        setf("call_attempts", n)
    n = _num(row.get("outlets"), int, "No. of outlets", problems)
    if n is not None:
        setf("distribution_outlets", n)
    n = _num(row.get("price_kg"), float, "Current price/kg", problems)
    if n is not None:
        setf("price_per_kg", n)

    d = _date(row.get("call_date"), "Call date", problems)
    if d:
        setf("last_call_date", d)
    fu = _txt(row.get("follow_up"))
    if fu:
        setf("next_followup_date", fu)

    if _txt(row.get("sample_req")):
        setf("sample_requested", _txt(row.get("sample_req")).lower() == "yes")
    et = _txt(row.get("email_type"))
    if et:
        setf("email_is_generic", et.lower() == "generic")
    if _txt(row.get("has_wa")).lower() == "yes" and not _txt(lead.whatsapp_number):
        setf("whatsapp_number", _txt(lead.phone))
    if _txt(row.get("is_dm")).lower() == "yes":
        setf("decision_maker", _txt(row.get("dm_name")) or _txt(row.get("spoke_with")))
    elif _txt(row.get("dm_name")):
        setf("decision_maker", _txt(row.get("dm_name")))

    consent = _txt(row.get("consent")).lower()
    if consent == "yes":
        setf("consent_status", CONSENT_YES)
        setf("consent_source", "FOUNDER_CALL")
        setf("consent_timestamp", d or datetime.utcnow())
    elif consent == "no":
        setf("consent_status", CONSENT_NO)

    if _txt(row.get("outcome")) == "Do not call again":
        setf("do_not_call", True)
        setf("dnc_reason", "asked on founder call "
                           + (d.strftime("%Y-%m-%d") if d else ""))
    return ch


def apply_row(row, lead, ch, db, problems):
    """Write, routing anything gated through the service that owns the gate."""
    from app.models.models import CallHistory, ObjectionLearning
    from app.services import trust_promoter as tp

    notes = []
    for f, v in ch.items():
        setattr(lead, f, v)

    email = _txt(row.get("email_new"))
    if email and "@" in email:
        if _txt(lead.email).lower() != email.lower():
            lead.email = email
        lead.email_source = "FOUNDER_CALL"
        # The gate's own entry point: promotes to VERIFIED and recomputes
        # confidence. Setting email_trust here by hand would skip the
        # shared-inbox cap and the audit trail.
        state = tp.on_founder_call(lead, db, confirmed=True,
                                   notes=_txt(row.get("remarks"))[:300])
        notes.append(f"trust -> {state}")
    elif email:
        problems.append(f"EMAIL captured: {email!r} is not an address — skipped")

    d = _date(row.get("call_date"), "Call date", [])
    if _txt(row.get("outcome")):
        db.add(CallHistory(
            lead_id=lead.id,
            call_date=d or datetime.utcnow(),
            status=_txt(row.get("outcome")),
            summary=_txt(row.get("remarks"))[:1000] or None,
        ))
        notes.append("call logged")

    why_stay, why_switch = _txt(row.get("why_stay")), _txt(row.get("why_switch"))
    competitor = _txt(row.get("competitor")) or _txt(row.get("brand_now"))
    if why_stay or why_switch or _txt(row.get("objection")):
        body = "  |  ".join(x for x in (
            f"STAYS BECAUSE: {why_stay}" if why_stay else "",
            f"WOULD SWITCH IF: {why_switch}" if why_switch else "",
        ) if x)
        db.add(ObjectionLearning(
            lead_id=lead.id,
            objection_type=_txt(row.get("objection")) or "call_intel",
            competitor_name=competitor or None,
            offered_price_per_kg=_num(row.get("price_kg"), float, "price", []),
            recorded_at=d or datetime.utcnow(),
            notes=body or None,
        ))
        notes.append("learning captured")
    return notes


def main():
    ap = argparse.ArgumentParser(description="Import a filled call sheet.")
    ap.add_argument("file")
    ap.add_argument("--commit", action="store_true", help="actually write")
    args = ap.parse_args()

    if not os.path.exists(args.file):
        raise SystemExit(f"No such file: {args.file}")

    from app.database.database import SessionLocal
    from app.models.models import B2BLead

    rows = read_rows(args.file)
    if not rows:
        print("No filled rows found. Nothing to import.")
        return 0

    db = SessionLocal()
    matched = unmatched = changed = 0
    problems, log = [], []
    try:
        by_phone = {}
        for l in db.query(B2BLead).filter(B2BLead.phone.isnot(None)).all():
            by_phone.setdefault(_digits(l.phone), l)

        for row in rows:
            lead = by_phone.get(_digits(row.get("phone")))
            if not lead:
                unmatched += 1
                problems.append(f"row {row['_row']}: no lead with phone "
                                f"{_txt(row.get('phone'))!r} — skipped")
                continue
            matched += 1
            rp = []
            ch = plan(row, lead, rp)
            notes = apply_row(row, lead, ch, db, rp) if args.commit else []
            if ch or notes:
                changed += 1
                log.append((lead.company, len(ch), notes))
            problems.extend(f"row {row['_row']}: {p}" for p in rp)

        if args.commit:
            db.commit()
    finally:
        db.close()

    print(f"rows with entries : {len(rows)}")
    print(f"matched on phone  : {matched}")
    print(f"unmatched         : {unmatched}")
    print(f"would change      : {changed}" if not args.commit else f"changed: {changed}")
    for company, n, notes in log[:25]:
        print(f"   {str(company)[:38]:<40}{n} field(s)  {', '.join(notes)}")
    if len(log) > 25:
        print(f"   ... and {len(log) - 25} more")
    if problems:
        print(f"\nskipped, nothing guessed ({len(problems)}):")
        for p in problems[:25]:
            print(f"   {p}")
        if len(problems) > 25:
            print(f"   ... and {len(problems) - 25} more")
    if not args.commit:
        print("\nDRY RUN — nothing written. Re-run with --commit to apply.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
