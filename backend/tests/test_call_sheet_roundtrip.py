"""
The call sheet has to close the loop, or it is just a spreadsheet.

A founder call is the only event that clears both send gates at once: it
produces first-party evidence (trust -> VERIFIED) AND the interaction that
confidence is computed from. These tests prove the import path actually
delivers that, and that it refuses the three things a data-capture layer must
never do — overwrite with blanks, guess at unreadable values, or manufacture
trust by writing email_trust directly.
"""
from __future__ import annotations

import os
import sys
import tempfile

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts")))

from app.models.models import (Base, B2BLead, CallHistory, ObjectionLearning,
                               WorkflowEvent)
from app.services import trust_promoter as tp

import call_sheet_schema as schema
import import_call_sheet as imp


@pytest.fixture
def db():
    fd, path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    eng = create_engine(f"sqlite:///{path}")
    Base.metadata.create_all(eng)
    s = sessionmaker(bind=eng)()
    yield s
    s.close()
    eng.dispose()
    try:
        os.unlink(path)
    except OSError:
        pass


def _sheet(rows: list[dict], path: str):
    """Write a workbook shaped exactly like the exporter's CALL QUEUE."""
    wb = Workbook()
    ws = wb.active
    ws.title = imp.SHEET
    for i, (_k, title, *_rest) in enumerate(schema.COLUMNS, start=1):
        ws.cell(imp.HEADER_ROW, i, title)
    for r, data in enumerate(rows, start=imp.HEADER_ROW + 1):
        for i, (key, *_rest) in enumerate(schema.COLUMNS, start=1):
            if key in data:
                ws.cell(r, i, data[key])
    wb.save(path)
    return path


@pytest.fixture
def xlsx():
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    yield path
    try:
        os.unlink(path)
    except OSError:
        pass


def _lead(db, **kw):
    lead = B2BLead(company=kw.pop("company", "Abohar Coffee House"),
                   phone=kw.pop("phone", "+91-98765-43210"),
                   city="Abohar", division="HORECA", **kw)
    db.add(lead)
    db.commit()
    return lead


def _run(path, db):
    rows = imp.read_rows(path)
    by_phone = {imp._digits(l.phone): l for l in db.query(B2BLead).all()}
    problems, applied = [], 0
    for row in rows:
        lead = by_phone.get(imp._digits(row.get("phone")))
        if not lead:
            continue
        ch = imp.plan(row, lead, problems)
        imp.apply_row(row, lead, ch, db, problems)
        applied += 1
    db.commit()
    return applied, problems


def test_a_call_clears_both_send_gates(db, xlsx):
    """The whole point. Before: unsendable. After one call: sendable."""
    lead = _lead(db, email="", email_trust="UNSEEN")
    assert tp.may_send(lead)[0] is False

    _sheet([{
        "phone": "+91-98765-43210",
        "call_date": "2026-08-22",
        "outcome": "Connected",
        "spoke_with": "Store owner",
        "their_role": "Owner",
        "email_new": "owner@abohorcoffeehouse.in",
        "consent": "Yes",
        "is_dm": "Yes",
        "remarks": "Buys 6kg a month, wants to taste first.",
    }], xlsx)
    _run(xlsx, db)
    db.refresh(lead)

    assert lead.email == "owner@abohorcoffeehouse.in"
    assert lead.email_source == "FOUNDER_CALL"
    assert lead.email_trust == "VERIFIED", f"trust stuck at {lead.email_trust}"
    assert (lead.email_confidence or 0) >= tp.CONFIDENCE_FLOOR, (
        f"confidence {lead.email_confidence} still under the floor — the call "
        f"did not clear gate 2")

    ok, why = tp.may_send(lead)
    assert ok is True, f"still unsendable after a confirmed call: {why}"
    assert lead.consent_status == imp.CONSENT_YES


def test_trust_moves_through_the_engine_not_by_hand(db, xlsx):
    """If the importer wrote email_trust directly it would bypass the
    shared-inbox cap and leave no audit trail. Assert the transition was
    recorded by the trust engine itself."""
    lead = _lead(db, email="", email_trust="UNSEEN")
    _sheet([{"phone": "+91-98765-43210", "outcome": "Connected",
             "email_new": "owner@abohorcoffeehouse.in"}], xlsx)
    _run(xlsx, db)

    evs = db.query(WorkflowEvent).filter(
        WorkflowEvent.lead_id == lead.id,
        WorkflowEvent.event_type == "TRUST_TRANSITION").all()
    assert evs, "no TRUST_TRANSITION written — trust was set behind the engine's back"
    assert any((e.payload or {}).get("source") == "FOUNDER_CALL" for e in evs)


def test_blank_never_overwrites(db, xlsx):
    """An empty cell means 'did not ask', not 'no'. Losing a known supplier
    because a caller skipped the column would be worse than not importing."""
    lead = _lead(db, current_brand="Nescafe", current_supplier="Ludhiana Traders",
                 contact_name="Existing Contact")
    _sheet([{"phone": "+91-98765-43210", "outcome": "No answer"}], xlsx)
    _run(xlsx, db)
    db.refresh(lead)

    assert lead.current_brand == "Nescafe"
    assert lead.current_supplier == "Ludhiana Traders"
    assert lead.contact_name == "Existing Contact"


def test_unreadable_numbers_are_skipped_not_guessed(db, xlsx):
    """'8-10' is a range a human wrote. It is not 9, and it is not 8."""
    lead = _lead(db)
    # Capture the real starting values: distribution_outlets defaults to 0, not
    # None, so "is None" would test the column default rather than the
    # importer. Unchanged is the actual contract.
    before = (lead.price_per_kg, lead.distribution_outlets)

    _sheet([{"phone": "+91-98765-43210", "outcome": "Connected",
             "kg_month": "8-10", "price_kg": "about four hundred",
             "outlets": "a few"}], xlsx)
    _, problems = _run(xlsx, db)
    db.refresh(lead)

    assert (lead.price_per_kg, lead.distribution_outlets) == before, (
        f"a value was invented from unreadable text: {before} -> "
        f"{(lead.price_per_kg, lead.distribution_outlets)}")
    # free-text column keeps the human's words verbatim
    assert lead.monthly_consumption == "8-10"
    assert any("price" in p.lower() for p in problems), problems
    assert any("outlet" in p.lower() for p in problems), problems


def test_their_words_are_captured_as_learning(db, xlsx):
    """'Won't change, delivery is reliable' has to survive as text. A tick-box
    reduces it to 'Not interested', which teaches the messaging engine nothing."""
    lead = _lead(db)
    _sheet([{
        "phone": "+91-98765-43210", "outcome": "Connected",
        "objection": "Happy with current", "competitor": "Nescafe",
        "why_stay": "Delivery is reliable and they give 30 day credit",
        "why_switch": "Only if someone matches the credit terms",
    }], xlsx)
    _run(xlsx, db)

    ol = db.query(ObjectionLearning).filter(
        ObjectionLearning.lead_id == lead.id).one()
    assert ol.competitor_name == "Nescafe"
    assert "30 day credit" in ol.notes
    assert "WOULD SWITCH IF" in ol.notes
    assert db.query(CallHistory).filter(CallHistory.lead_id == lead.id).count() == 1


def test_do_not_call_is_honoured(db, xlsx):
    lead = _lead(db)
    _sheet([{"phone": "+91-98765-43210", "outcome": "Do not call again",
             "call_date": "2026-08-22"}], xlsx)
    _run(xlsx, db)
    db.refresh(lead)

    assert lead.do_not_call is True
    assert "founder call" in (lead.dnc_reason or "")


def test_exporter_and_importer_cannot_drift(db):
    """Both read the same COLUMNS. If someone re-declares them in one file, a
    header rename stops matching and the import silently writes nothing."""
    import export_call_sheet as exp

    assert exp.COLUMNS is schema.COLUMNS
    assert imp.COLUMNS is schema.COLUMNS
    targets = [c[5] for c in schema.COLUMNS if c[5] and ":" not in c[5]]
    cols = {c.name for c in B2BLead.__table__.columns}
    unknown = [t for t in targets if t not in cols]
    assert not unknown, f"columns target fields that do not exist: {unknown}"
